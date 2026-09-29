"""Render a FinalResume as Markdown or Word. Single column, plain structure: reads well in
applicant tracking systems."""

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from lodestar.resume.assemble import FinalResume

ACCENT = RGBColor(0x24, 0x47, 0x6B)


def _contact_line(resume: FinalResume) -> str:
    c = resume.contact
    if not c:
        return ""
    return " | ".join(x for x in [c.location, c.email, c.phone, *c.links.values()] if x)


def to_markdown(resume: FinalResume) -> str:
    lines = [f"# {resume.name}"]
    if contact := _contact_line(resume):
        lines.append(contact)
    lines += ["", f"**{resume.headline}**", "", resume.summary, "", "## Skills", "", ", ".join(resume.skills),
              "", "## Experience", ""]
    for role in resume.experience:
        lines.append(f"### {role.title}, {role.organization}")
        lines.append(f"*{role.dates}*")
        lines += [f"- {h}" for h in role.highlights]
        lines.append("")
    if resume.projects:
        lines += ["## Projects", ""]
        for p in resume.projects:
            lines.append(f"### {p.name}" + (f" ({p.url})" if p.url else ""))
            lines += [f"- {h}" for h in p.highlights]
            lines.append("")
    if resume.education:
        lines += ["## Education", ""]
        lines += [f"- {e.title}, {e.institution}" + (f", {e.year}" if e.year else "") for e in resume.education]
    return "\n".join(lines).rstrip() + "\n"


def _rule(paragraph) -> None:
    """A bottom border under a paragraph (a horizontal rule without a table)."""
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "24476B")):
        bottom.set(qn(k), v)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text.upper())
    run.bold, run.font.size, run.font.color.rgb = True, Pt(10.5), ACCENT
    _rule(p)


def to_docx(resume: FinalResume) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)       # US Letter
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(0.7))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(2)

    name = doc.add_paragraph()
    run = name.add_run(resume.name)
    run.bold, run.font.size = True, Pt(18)
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if contact := _contact_line(resume):
        doc.add_paragraph(contact).runs[0].font.size = Pt(9.5)
    headline = doc.add_paragraph()
    headline.add_run(resume.headline).bold = True
    doc.add_paragraph(resume.summary)

    _heading(doc, "Skills")
    doc.add_paragraph(", ".join(resume.skills))

    _heading(doc, "Experience")
    for role in resume.experience:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.add_run(f"{role.title}, {role.organization}").bold = True
        p.add_run(f"    {role.dates}").italic = True
        for h in role.highlights:
            doc.add_paragraph(h, style="List Bullet")

    if resume.projects:
        _heading(doc, "Projects")
        for proj in resume.projects:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(5)
            p.add_run(proj.name).bold = True
            if proj.url:
                p.add_run(f"    {proj.url}")
            for h in proj.highlights:
                doc.add_paragraph(h, style="List Bullet")

    if resume.education:
        _heading(doc, "Education")
        for e in resume.education:
            doc.add_paragraph(f"{e.title}, {e.institution}" + (f", {e.year}" if e.year else ""))

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
