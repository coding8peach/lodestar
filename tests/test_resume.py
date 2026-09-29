"""Resume tailoring: Python's checks, assembly, rendering, the resume agent (fake models, real MCP
server subprocess), the A2A round trip, and the service/CLI path."""

import asyncio
import io
import json
from contextlib import closing
from datetime import datetime, timezone

import httpx
import pytest
from conftest import EXAMPLE_PROFILE
from docx import Document
from fakes import FakeLlm, RateLimited, text
from google.genai import types

from lodestar.db import connect, get_status, init_db, record_decision, save_fit_result, save_job
from lodestar.resume.assemble import assemble
from lodestar.resume.render import to_docx, to_markdown
from lodestar.resume.validate import ResumeRejected, add_required_entries, check_resume, validate_resume
from lodestar.schemas import FitAnalysis, JobPosting, RequirementMatch, job_id_from_url, load_profile
from lodestar.schemas.resume import TailoredResume
from lodestar.scoring import score_analysis

PROFILE = load_profile(EXAMPLE_PROFILE)
URL = "https://boards.greenhouse.io/acme/jobs/555"
JOB_ID = job_id_from_url(URL)


def draft(**overrides) -> TailoredResume:
    data = {
        "job_id": JOB_ID,
        "headline": "Backend engineer: Java and Python",
        "summary": {"text": "Backend engineer with 10+ years of Java on Oracle.", "sources": ["exp-acme"]},
        "skills": ["Java", "Python"],
        "experience": [{"entry_id": "exp-acme", "highlights": [
            {"text": "Built the data pipeline for materials experiments in Java", "sources": ["exp-acme"]}]}],
        "projects": [{"entry_id": "proj-quizbot", "highlights": [
            {"text": "Built an MCP-based quiz emailer in Python", "sources": ["proj-quizbot"]}]}],
        "notes": ["Kubernetes not claimed: not in the profile."],
    }
    data.update(overrides)
    return TailoredResume.model_validate(data)


# --- Python's checks ------------------------------------------------------------------------

def test_valid_draft_passes():
    assert validate_resume(draft(), PROFILE, JOB_ID) == []


def test_invented_skill_rejected():
    problems = validate_resume(draft(skills=["Java", "Kubernetes"]), PROFILE, JOB_ID)
    assert problems == ["skills not in the profile: Kubernetes"]


def test_unknown_entry_and_source_rejected():
    bad = draft(experience=[{"entry_id": "exp-google", "highlights": [
        {"text": "Led search infrastructure", "sources": ["exp-google"]}]}])
    problems = validate_resume(bad, PROFILE, JOB_ID)
    assert any("'exp-google' is not a experience entry" in p for p in problems)
    assert any("unknown source(s) exp-google" in p for p in problems)


def test_invented_number_rejected_but_career_years_allowed():
    bad = draft(experience=[{"entry_id": "exp-acme", "highlights": [
        {"text": "Cut pipeline runtime by 40%", "sources": ["exp-acme"]}]}])
    assert any("number(s) 40 not in its sources" in p for p in validate_resume(bad, PROFILE, JOB_ID))
    ok = draft(summary={"text": "14+ years of backend engineering.", "sources": ["exp-acme"]})
    assert validate_resume(ok, PROFILE, JOB_ID) == []  # career starts 2012: 14 years is supported


def test_line_without_sources_rejected():
    bad = draft(summary={"text": "Great engineer.", "sources": []})
    assert any("cites no profile entry" in p for p in validate_resume(bad, PROFILE, JOB_ID))


def test_check_resume_raises_with_every_problem():
    with pytest.raises(ResumeRejected, match="Kubernetes"):
        check_resume(draft(skills=["Kubernetes"]), PROFILE, JOB_ID)


def test_recent_roles_are_added_if_left_out():
    result = add_required_entries(draft(experience=[]), PROFILE)
    assert [e.entry_id for e in result.experience] == ["exp-acme"]  # 2012-2019: since 2006, always kept
    assert any("always kept" in n for n in result.notes)


# --- assembly and rendering -------------------------------------------------------------------

def test_assemble_fills_facts_from_the_profile():
    final = assemble(draft(), PROFILE)
    role = final.experience[0]
    assert (role.title, role.organization, role.dates) == ("Staff Software Engineer", "Acme Corp", "2012 – 2019")
    assert final.name == PROFILE.name and final.contact.email == "jane@example.com"
    assert [e.title for e in final.education][0] == "Agent frameworks (ADK, A2A, MCP)"  # newest first


def test_markdown_and_docx():
    final = assemble(draft(), PROFILE)
    md = to_markdown(final)
    assert md.startswith(f"# {PROFILE.name}") and "## Experience" in md and "- Built the data pipeline" in md
    doc = Document(io.BytesIO(to_docx(final)))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == PROFILE.name and "EXPERIENCE" in texts
    assert any(p.style.name == "List Bullet" for p in doc.paragraphs)
    assert round(doc.sections[0].page_width.inches, 1) == 8.5  # US Letter


# --- the resume agent: fake models, real MCP server -------------------------------------------

@pytest.fixture
def data(tmp_path, monkeypatch):
    import shutil
    db, profile = tmp_path / "t.sqlite", tmp_path / "profile.yaml"
    shutil.copy(EXAMPLE_PROFILE, profile)
    monkeypatch.setenv("LODESTAR_DB", str(db))
    monkeypatch.setenv("LODESTAR_PROFILE", str(profile))
    with closing(connect(db)) as conn:
        init_db(conn)
        save_job(conn, JobPosting(id=JOB_ID, url=URL, source="greenhouse", classified_by="url_rule",
                                  title="Senior Backend Engineer", company="Beta", description="Java. " * 60,
                                  fetched_at=datetime(2026, 9, 27, tzinfo=timezone.utc)))
        save_fit_result(conn, score_analysis(FitAnalysis(job_id=JOB_ID, summary="s", requirements=[
            RequirementMatch(requirement="Java", priority="must_have", evidence=["exp-acme"],
                             match_level="direct", explanation="")])), model="m")
        record_decision(conn, JOB_ID, "approve")
    return db


TOOL_CALLS = [
    types.Part(function_call=types.FunctionCall(name="get_job", args={"job_id": JOB_ID})),
    types.Part(function_call=types.FunctionCall(name="get_profile", args={})),
    types.Part(function_call=types.FunctionCall(name="get_fit_analysis", args={"job_id": JOB_ID})),
]
GOOD = draft().model_dump_json()
BAD = draft(skills=["Java", "Kubernetes"]).model_dump_json()


def test_resume_agent_reads_three_tools_and_retries_a_rejected_draft(data):
    from lodestar.resume_agent.runner import tailor_job

    fake = FakeLlm(model="a", script=[TOOL_CALLS, text(BAD), text(GOOD)])
    run = asyncio.run(tailor_job(JOB_ID, ["a"], PROFILE, make_model=lambda n: fake))
    assert run.result.skills == ["Java", "Python"] and run.model == "a"
    tool_results = {p.function_response.name: p.function_response.response
                    for c in fake.requests[1].contents for p in c.parts if p.function_response}
    assert set(tool_results) == {"get_job", "get_profile", "get_fit_analysis"}
    assert tool_results["get_fit_analysis"]["recommendation"] == "top_pick"
    assert "Kubernetes" in fake.requests[2].contents[-1].parts[0].text  # the retry says what was wrong


def test_resume_agent_uses_the_shared_fallback(data):
    from lodestar.resume_agent.runner import tailor_job

    a = FakeLlm(model="a", script=[RateLimited()])
    b = FakeLlm(model="b", script=[TOOL_CALLS, text(GOOD)])
    run = asyncio.run(tailor_job(JOB_ID, ["a", "b"], PROFILE, make_model={"a": a, "b": b}.get))
    assert run.model == "b" and run.skipped[0].startswith("a: RateLimited")


def test_resume_agent_over_a2a(data):
    from lodestar.fit_agent.client import tailor_via_a2a
    from lodestar.fit_agent.server import build_resume_app

    fake = FakeLlm(model="a", script=[TOOL_CALLS, text(GOOD)])
    app = build_resume_app(["a"], port=8002, make_model=lambda n: fake)

    async def go():
        async with app.router.lifespan_context(app):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://localhost:8002",
                                         timeout=60) as http:
                card = (await http.get("/.well-known/agent-card.json")).json()
                reply = await tailor_via_a2a(JOB_ID, base_url="http://localhost:8002", httpx_client=http)
                return card, reply

    card, reply = asyncio.run(go())
    assert card["name"] == "resume_agent"
    assert reply.resume.headline == "Backend engineer: Java and Python" and reply.model == "a"


# --- service and CLI ----------------------------------------------------------------------------

def _reply(resume=None):
    from lodestar.schemas.resume import ResumeAgentReply

    async def fake(job_id, max_paid_usd=None):
        return ResumeAgentReply(resume=resume or draft(), model="gemini-test")
    return fake


def test_service_tailor_stores_and_moves_the_status(data):
    from lodestar.app import service

    view = service.tailor(JOB_ID, tailor_fn=_reply())
    assert view.markdown.startswith(f"# {PROFILE.name}") and view.filename == "beta-senior-backend-engineer"
    with closing(connect()) as conn:
        assert get_status(conn, JOB_ID) == "resume_tailored"
        assert conn.execute("SELECT COUNT(*) FROM resumes").fetchone()[0] == 1
    service.tailor(JOB_ID, tailor_fn=_reply())  # tailoring again: a new version, status unchanged
    with closing(connect()) as conn:
        assert conn.execute("SELECT COUNT(*) FROM resumes").fetchone()[0] == 2
    assert service.resume_docx(JOB_ID)[:2] == b"PK"  # a zip, i.e. a .docx


def test_service_rechecks_the_draft_locally(data):
    from lodestar.app import service

    with pytest.raises(ValueError, match="didn't hold up.*Kubernetes"):
        service.tailor(JOB_ID, tailor_fn=_reply(draft(skills=["Kubernetes"])))


def test_only_approved_jobs(data):
    from lodestar.app import service

    with pytest.raises(ValueError, match="only approved jobs"):
        with closing(connect()) as conn:
            conn.execute("UPDATE jobs SET status = 'rejected'")
            conn.commit()
        service.tailor(JOB_ID, tailor_fn=_reply())


def test_cli_tailor_writes_files(data, monkeypatch, tmp_path):
    from lodestar.app import service
    from lodestar.workflow import cli

    real = service.tailor
    monkeypatch.setattr(service, "tailor", lambda job_id: real(job_id, tailor_fn=_reply()))
    monkeypatch.setattr(cli, "DATA_DIR", tmp_path)
    assert cli.tailor_cmd([JOB_ID]) == 0
    assert (tmp_path / "resumes" / "beta-senior-backend-engineer.md").exists()
    assert (tmp_path / "resumes" / "beta-senior-backend-engineer.docx").exists()


def test_get_fit_analysis_tool(data):
    from mcp import Client

    from lodestar.mcp_server.server import mcp

    async def go():
        async with Client(mcp) as client:
            return (await client.call_tool("get_fit_analysis", {"job_id": JOB_ID}),
                    await client.call_tool("get_fit_analysis", {"job_id": "job_0000000000000000"}))

    found, missing = asyncio.run(go())
    assert found.structured_content["recommendation"] == "top_pick"
    assert missing.is_error and "no fit analysis" in json.dumps([c.text for c in missing.content])


# --- impact words and placeholders (2b) -----------------------------------------------------

def test_impact_words_need_support():
    puffed = draft(summary={"text": "Seasoned engineer with a proven track record of robust APIs.",
                            "sources": ["exp-acme"]})
    problems = validate_resume(puffed, PROFILE, JOB_ID)
    assert len(problems) == 1 and all(w in problems[0] for w in ("proven", "track record", "robust", "seasoned"))


def test_impact_word_allowed_when_the_source_uses_it():
    profile = PROFILE.model_copy(update={"summary": "Seasoned backend engineer."})
    ok = draft(summary={"text": "Seasoned backend engineer working in Java.", "sources": ["exp-acme"]})
    assert validate_resume(ok, profile, JOB_ID) == []


def test_placeholder_in_agent_text_is_rejected():
    bad = draft(headline="Backend engineer at <company>")
    assert any("placeholder" in p for p in validate_resume(bad, PROFILE, JOB_ID))


def test_placeholders_in_profile_facts_are_found():
    from lodestar.resume.validate import profile_placeholders

    edu = [e.model_copy(update={"institution": "TODO course provider"}) if e.id == "course-agents" else e
           for e in PROFILE.education]
    profile = PROFILE.model_copy(update={"education": edu})
    assert profile_placeholders(draft(), profile) == ["course-agents: 'TODO course provider'"]
    assert profile_placeholders(draft(), PROFILE) == []


def test_service_refuses_to_save_with_profile_placeholders(data):
    import yaml

    from lodestar.app import service

    path = data.parent / "profile.yaml"
    cfg = yaml.safe_load(path.read_text())
    for e in cfg["education"]:
        if e["id"] == "course-agents":
            e["institution"] = "TODO course provider"
    path.write_text(yaml.safe_dump(cfg))
    with pytest.raises(ValueError, match="placeholders.*TODO course provider"):
        service.tailor(JOB_ID, tailor_fn=_reply())
    with closing(connect()) as conn:
        assert conn.execute("SELECT COUNT(*) FROM resumes").fetchone()[0] == 0


# --- the Decisions page's resume panel ------------------------------------------------------------

def _panel_app():
    def script():
        from lodestar.app import service as svc
        from lodestar.ui import pages
        pages._resume_panel(svc.decisions("approve")[0])

    from streamlit.testing.v1 import AppTest
    at = AppTest.from_function(script, default_timeout=30)
    at.run()
    assert not at.exception, at.exception
    return at


def test_panel_tailors_and_offers_downloads(data, monkeypatch):
    from lodestar.app import service
    from lodestar.fit_agent import client

    monkeypatch.setattr(client, "agent_status", lambda url=None: (True, "running"))
    real = service.tailor
    monkeypatch.setattr(service, "tailor", lambda job_id: real(job_id, tailor_fn=_reply()))
    at = _panel_app()
    assert at.button(key="tailor").label == "Tailor resume" and not at.button(key="tailor").disabled
    assert "No resume yet" in at.info[0].value
    at.button(key="tailor").click().run()
    assert not at.exception
    assert service.decisions("approve")[0].has_resume
    at = _panel_app()
    assert at.button(key="tailor").label == "Tailor again"
    assert any(PROFILE.name in m.value for m in at.markdown)  # the preview


def test_panel_disabled_when_the_agent_is_down(data, monkeypatch):
    from lodestar.fit_agent import client

    monkeypatch.setattr(client, "agent_status", lambda url=None: (False, "not running"))
    at = _panel_app()
    assert at.button(key="tailor").disabled
    assert any("lodestar-agent" in c.value for c in at.caption)


def test_calls_recorded_even_when_the_local_check_stops_the_save(data):
    from lodestar.app import service
    from lodestar.schemas import LlmCall
    from lodestar.schemas.resume import ResumeAgentReply

    async def reply_with_calls(job_id, max_paid_usd=None):
        return ResumeAgentReply(resume=draft(skills=["Kubernetes"]), model="m", calls=[LlmCall(
            model="m", attempt=1, status="ok", input_tokens=100, started_at=datetime.now(timezone.utc))])

    with pytest.raises(ValueError, match="didn't hold up"):
        service.tailor(JOB_ID, tailor_fn=reply_with_calls)
    with closing(connect()) as conn:
        assert conn.execute("SELECT COUNT(*) FROM llm_calls").fetchone()[0] == 1
        assert conn.execute("SELECT COUNT(*) FROM resumes").fetchone()[0] == 0
